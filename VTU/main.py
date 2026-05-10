import time
import pandas as pd
import re
import cv2
from PIL import Image
import easyocr

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import NoAlertPresentException

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment

# =============================
# LOAD USN LIST
# =============================

df = pd.read_excel("VTU/usn_list.xlsx")
usn_list = df["USN"].tolist()

# =============================
# OCR INIT
# =============================

reader = easyocr.Reader(['en'], gpu=False)

def preprocess_image(path):
    image = cv2.imread(path)

    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
    blur = cv2.GaussianBlur(gray, (5,5), 0)

    _, thresh = cv2.threshold(
        blur,
        0,
        255,
        cv2.THRESH_BINARY + cv2.THRESH_OTSU
    )

    return thresh

# =============================
# BROWSER
# =============================

driver = webdriver.Chrome()
wait = WebDriverWait(driver, 15)

driver.get("https://results.vtu.ac.in/D25J26Ecbcs/index.php")

all_students = []

# =============================
# CAPTCHA SOLVER FUNCTION
# =============================

def solve_captcha(usn):

    while True:
        try:
            # -------------------------
            # ALWAYS re-enter USN
            # -------------------------
            usn_box = wait.until(EC.presence_of_element_located((By.NAME, "lns")))
            usn_box.clear()
            usn_box.send_keys(usn)

            # -------------------------
            # locate captcha
            # -------------------------
            captcha = wait.until(
                EC.presence_of_element_located((By.XPATH, "//img[contains(@src,'captcha')]"))
            )

            # -------------------------
            # crop captcha
            # -------------------------
            location = captcha.location
            size = captcha.size

            driver.save_screenshot("full.png")
            img = Image.open("full.png")

            left = location['x']
            top = location['y']
            right = left + size['width']
            bottom = top + size['height']

            captcha_img = img.crop((left, top, right, bottom))
            captcha_img.save("temp.png")

            # -------------------------
            # preprocess
            # -------------------------
            processed = preprocess_image("temp.png")

            # -------------------------
            # OCR
            # -------------------------
            results = reader.readtext(processed)

            detected_text = ""
            for r in results:
                detected_text += r[1]

            prediction = detected_text.strip()

            # -------------------------
            # clean
            # -------------------------
            prediction = re.sub(r'[^A-Za-z0-9]', '', prediction)

            if len(prediction) != 6:
                print("Invalid OCR, retrying...")
                driver.refresh()
                time.sleep(1)
                continue

            print(f"Captcha attempt: {prediction}")

            # -------------------------
            # enter captcha
            # -------------------------
            captcha_box = driver.find_element(By.NAME, "captchacode")
            captcha_box.clear()
            captcha_box.send_keys(prediction)

            # -------------------------
            # submit
            # -------------------------
            driver.find_element(By.ID, "submit").click()

            time.sleep(1)

            # -------------------------
            # check failure
            # -------------------------
            try:
                alert = driver.switch_to.alert
                if "Invalid captcha" in alert.text:
                    alert.accept()
                    print("❌ Wrong captcha")
                    time.sleep(1)
                    continue

            except NoAlertPresentException:
                pass

            # -------------------------
            # success check
            # -------------------------
            wait.until(EC.presence_of_element_located((By.CLASS_NAME, "divTableRow")))

            print("✅ Captcha solved")
            return True

        except Exception as e:
            print("Captcha error:", e)
            driver.refresh()
            time.sleep(1)

# =============================
# MAIN LOOP
# =============================

for i, usn in enumerate(usn_list, 1):

    print(f"\nProcessing {i}/{len(usn_list)} : {usn}")

    success = False

    while not success:
        try:
            # enter USN
            usn_box = wait.until(EC.presence_of_element_located((By.NAME, "lns")))
            usn_box.clear()
            usn_box.send_keys(usn)

            # solve captcha
            success = solve_captcha(usn)

        except Exception as e:
            print("Retrying USN due to error:", e)
            driver.get("https://results.vtu.ac.in/D25J26Ecbcs/index.php")
            time.sleep(2)

    # =============================
    # SCRAPE DATA (UNCHANGED)
    # =============================

    soup = BeautifulSoup(driver.page_source, "html.parser")

    student_data = {}

    try:
        rows = soup.find_all("tr")

        for r in rows:
            if "Student Name" in r.text:
                tds = r.find_all("td")
                student_data["Name"] = tds[1].text.replace(":", "").strip()

            if "University Seat Number" in r.text:
                tds = r.find_all("td")
                student_data["USN"] = tds[1].text.replace(":", "").strip()

        subject_rows = soup.find_all("div", class_="divTableRow")[1:7]

        core_subject_codes = ["BAI701", "BAI702", "BAD703", "BAI786"]

        for idx, row in enumerate(subject_rows):

            cells = row.find_all("div", class_="divTableCell")

            subject_code = cells[0].text.strip()
            subject_name = cells[1].text.strip()
            internal = cells[2].text.strip()
            external = cells[3].text.strip()
            total = cells[4].text.strip()
            result = cells[5].text.strip()

            if idx < 4:
                sub = core_subject_codes[idx]

                student_data[f"{sub}_Internal"] = internal
                student_data[f"{sub}_External"] = external
                student_data[f"{sub}_Total"] = total
                student_data[f"{sub}_Result"] = result

            else:
                if subject_code.startswith("BAD714C"):
                    student_data["PEC_Subject"] = subject_name
                    student_data["PEC_Internal"] = internal
                    student_data["PEC_External"] = external
                    student_data["PEC_Total"] = total
                    student_data["PEC_Result"] = result
                else:
                    student_data["OEC_Subject"] = subject_name
                    student_data["OEC_Internal"] = internal
                    student_data["OEC_External"] = external
                    student_data["OEC_Total"] = total
                    student_data["OEC_Result"] = result

        all_students.append(student_data)

        print("🎯 Scraped successfully")

    except Exception as e:
        print("Scrape error:", e)

    # reload page
    driver.get("https://results.vtu.ac.in/D25J26Ecbcs/index.php")
    time.sleep(2)

driver.quit()

# =============================
# EXCEL (UNCHANGED)
# =============================

wb = Workbook()
ws = wb.active

ws.cell(row=1, column=1, value="Name")
ws.cell(row=1, column=2, value="USN")

col = 3
core_subjects = ["BAI701", "BAI702", "BAD703", "BAI786"]

for sub in core_subjects:
    ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col+3)
    ws.cell(row=1, column=col, value=sub)

    ws.cell(row=2, column=col, value="Internals")
    ws.cell(row=2, column=col+1, value="Externals")
    ws.cell(row=2, column=col+2, value="Total")
    ws.cell(row=2, column=col+3, value="Result")

    col += 4

ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col+4)
ws.cell(row=1, column=col, value="PEC")

ws.cell(row=2, column=col, value="Subject")
ws.cell(row=2, column=col+1, value="Internals")
ws.cell(row=2, column=col+2, value="Externals")
ws.cell(row=2, column=col+3, value="Total")
ws.cell(row=2, column=col+4, value="Result")

col += 5

ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col+4)
ws.cell(row=1, column=col, value="OEC")

ws.cell(row=2, column=col, value="Subject")
ws.cell(row=2, column=col+1, value="Internals")
ws.cell(row=2, column=col+2, value="Externals")
ws.cell(row=2, column=col+3, value="Total")
ws.cell(row=2, column=col+4, value="Result")

row_num = 3

for student in all_students:
    ws.cell(row=row_num, column=1, value=student.get("Name", ""))
    ws.cell(row=row_num, column=2, value=student.get("USN", ""))

    col = 3

    for sub in core_subjects:
        ws.cell(row=row_num, column=col, value=student.get(f"{sub}_Internal", ""))
        ws.cell(row=row_num, column=col+1, value=student.get(f"{sub}_External", ""))
        ws.cell(row=row_num, column=col+2, value=student.get(f"{sub}_Total", ""))
        ws.cell(row=row_num, column=col+3, value=student.get(f"{sub}_Result", ""))
        col += 4

    ws.cell(row=row_num, column=col, value=student.get("PEC_Subject", ""))
    ws.cell(row=row_num, column=col+1, value=student.get("PEC_Internal", ""))
    ws.cell(row=row_num, column=col+2, value=student.get("PEC_External", ""))
    ws.cell(row=row_num, column=col+3, value=student.get("PEC_Total", ""))
    ws.cell(row=row_num, column=col+4, value=student.get("PEC_Result", ""))

    col += 5

    ws.cell(row=row_num, column=col, value=student.get("OEC_Subject", ""))
    ws.cell(row=row_num, column=col+1, value=student.get("OEC_Internal", ""))
    ws.cell(row=row_num, column=col+2, value=student.get("OEC_External", ""))
    ws.cell(row=row_num, column=col+3, value=student.get("OEC_Total", ""))
    ws.cell(row=row_num, column=col+4, value=student.get("OEC_Result", ""))

    row_num += 1

for cell in ws[1]:
    cell.alignment = Alignment(horizontal="center")

wb.save("VTU/vtu_results.xlsx")

print("\nExcel file created: VTU/vtu_results.xlsx")
